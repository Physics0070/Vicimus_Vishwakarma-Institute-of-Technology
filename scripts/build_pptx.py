"""Build docs/GlucoTwin_Presentation.pptx: an 18-slide, 20-minute talk with a full speaker script in the notes.

Every number comes from reports/metrics.json, reports/real_data_metrics.json or the running API (via an
in-process test client), so the deck always matches the code. Requires the trained artifacts
(`python -m glucotwin.pipeline`) and python-pptx.

    python scripts/build_pptx.py
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
OUT = ROOT / "docs" / "GlucoTwin_Presentation.pptx"
IMG = ROOT / "docs" / "src"
FIG = ROOT / "reports" / "figures"

LIVE_URL = "https://glucotwin-india.vercel.app"
REPO_URL = "github.com/Physics0070/Vicimus_Vishwakarma-Institute-of-Technology"
REAL_DATA_STALE = True      # set False after re-running `python -m glucotwin.realdata` on the current code
DEMO_PID, DEMO_AT = "P050", "2026-09-11T12:10"

# ---------------------------------------------------------------------------- palette (the dashboard's)
INK, GRAPH = "15141A", "24222C"
IVORY, SAND = "F2EEE6", "B8B1A4"
WHITE, TINT, LINE = "FFFFFF", "F3F4FA", "DFE1EE"
TEXT, MUTED = "1E1D26", "5A5968"
PERI, ORCHID = "3F51D8", "A21CAF"
VERM, GOLD, GREEN = "D4351C", "B7791F", "15803D"
HEAD, BODY, MONO = "Cambria", "Calibri", "Courier New"

W, H = 13.333, 7.5
MX = 0.6                      # side margin
CW = W - 2 * MX               # content width


def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h)


# ---------------------------------------------------------------------------- facts
def facts() -> dict:
    m = json.loads((ROOT / "reports" / "metrics.json").read_text())
    r = json.loads((ROOT / "reports" / "real_data_metrics.json").read_text())
    f = m["final"]
    from fastapi.testclient import TestClient
    from index import app
    c = TestClient(app)
    d = c.get(f"/api/patients/{DEMO_PID}?at={DEMO_AT}").json()
    sc = [{"label": "a", "meals": [], "walks": []},
          {"label": "b", "meals": [{"meal_key": "rice_dal_sabzi", "portion": 1, "offset_min": 15}], "walks": []},
          {"label": "c", "meals": [{"meal_key": "rice_dal_sabzi", "portion": 1, "offset_min": 15}],
           "walks": [{"offset_min": 30, "duration_min": 15, "cadence": 100}]}]
    wi = c.post(f"/api/patients/{DEMO_PID}/whatif", json={"at": DEMO_AT, "horizon_min": 300, "scenarios": sc}).json()
    rev = c.get(f"/api/patients/{DEMO_PID}?at={DEMO_AT}&reveal=true").json()["outcome"]["cgm"]
    import joblib      # our own trained artifact, produced by `python -m glucotwin.pipeline` (same load the API does)
    groups = joblib.load(ROOT / "artifacts" / "models.joblib")["groups"]
    return {"m": m, "f": f, "r": r, "demo": d, "whatif": wi, "peak": max(v for v in rev if v is not None), "groups": {k: len(v) for k, v in groups.items()},
            "n_ward": len(m["split"]["test"]), "n_train": len(m["split"]["train"]), "n_cal": len(m["split"]["calib"])}


# ---------------------------------------------------------------------------- drawing helpers
def new_deck() -> Presentation:
    p = Presentation()
    p.slide_width, p.slide_height = Inches(W), Inches(H)
    return p


def add_slide(pres, title: str | None, dark: bool = False):
    s = pres.slides.add_slide(pres.slide_layouts[5])        # "Title Only": keeps a real title for outline/accessibility
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(INK if dark else WHITE)
    t = s.shapes.title
    if title is None:
        t._element.getparent().remove(t._element)
        return s
    t.left, t.top, t.width, t.height = Inches(MX), Inches(0.72), Inches(CW), Inches(0.75)
    tf = t.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.text = title
    para = tf.paragraphs[0]
    para.alignment = PP_ALIGN.LEFT
    r = para.runs[0]
    r.font.name, r.font.size, r.font.bold = HEAD, Pt(32), True
    r.font.color.rgb = rgb(IVORY if dark else TEXT)
    return s


def text(s, x, y, w, h, paras, size=14, color=TEXT, bold=False, font=BODY, align="l", anchor="t", name=None):
    """paras: str | list of (str | dict). dict keys: t / runs, size, bold, color, font, bullet, after, italic, align."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(paras, str):
        paras = [paras]
    for i, spec in enumerate(paras):
        spec = {"t": spec} if isinstance(spec, str) else dict(spec)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[spec.get("align", align)]
        runs = spec.get("runs") or [(spec["t"], {})]
        for rt, ro in runs:
            r = p.add_run()
            r.text = rt
            r.font.name = ro.get("font", spec.get("font", font))
            r.font.size = Pt(ro.get("size", spec.get("size", size)))
            r.font.bold = ro.get("bold", spec.get("bold", bold))
            r.font.italic = ro.get("italic", spec.get("italic", False))
            r.font.color.rgb = rgb(ro.get("color", spec.get("color", color)))
        if "after" in spec:
            p.space_after = Pt(spec["after"])
        if "before" in spec:
            p.space_before = Pt(spec["before"])
        if spec.get("bullet"):
            ind = Inches(0.22)
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(int(ind)))
            pPr.set("indent", str(-int(ind)))
            bu = etree.SubElement(pPr, qn("a:buChar"))
            bu.set("char", "•")
    return tb


def box(s, x, y, w, h, fill=TINT, line=None, radius=0.06, name=None):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = radius
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(fill)
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(0.75)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    if name:
        sh.name = name
    return sh


def dot(s, x, y, d, color):
    sh = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(color)
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def badge(s, x, y, d, label, fill, color=WHITE, size=16):
    sh = dot(s, x, y, d, fill)
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.name, r.font.size, r.font.bold = BODY, Pt(size), True
    r.font.color.rgb = rgb(color)
    return sh


def kicker(s, label, dark=False, color=None):
    """Series-key motif (observed / predicted / twin) + section label."""
    for i, c in enumerate((GRAPH if not dark else IVORY, PERI, ORCHID)):
        dot(s, MX + i * 0.17, 0.43, 0.11, c)
    text(s, MX + 0.62, 0.36, 8, 0.25, label.upper(), size=12, bold=True, color=color or (SAND if dark else MUTED))


def footer(s, n, source=None, dark=False):
    if source:
        text(s, MX, 6.78, CW, 0.25, source, size=11, color=SAND if dark else MUTED)
    text(s, MX, 7.08, 6, 0.22, "GlucoTwin India · Team Vicimus", size=10, color=SAND if dark else MUTED)
    text(s, W - MX - 1, 7.08, 1, 0.22, str(n), size=10, color=SAND if dark else MUTED, align="r")


def picture(s, path: Path, x, y, w=None, h=None, alt="", border=True):
    from PIL import Image
    iw, ih = Image.open(path).size
    ar = iw / ih
    if w and h:
        if w / h > ar:
            w = h * ar
        else:
            h = w / ar
    elif w:
        h = w / ar
    else:
        w = h * ar
    pic = s.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    pic._element.nvPicPr.cNvPr.set("descr", alt)
    if border:
        pic.line.color.rgb = rgb(LINE)
        pic.line.width = Pt(0.75)
    return pic, w, h


def notes(s, start, secs, script):
    mm = lambda t: f"{t // 60}:{t % 60:02d}"
    s.notes_slide.notes_text_frame.text = f"[{mm(start)} to {mm(start + secs)} · {secs} s]\n\n{script.strip()}"


def style_chart_text(chart, size=12, color=TEXT):
    chart.font.size = Pt(size)
    chart.font.name = BODY
    chart.font.color.rgb = rgb(color)


# ---------------------------------------------------------------------------- the deck
def build() -> None:
    F = facts()
    f, m, r, d, wi = F["f"], F["m"], F["r"], F["demo"], F["whatif"]
    pers = m["baselines"]["Persistence"]
    ev = f["events"]
    rd = r["ablation"]["+ Twin physiology (full hybrid)"]
    rp = r["baselines"]["Persistence"]
    re_ = r["events_180"]
    sc = r["status_counts"]
    TW = m["twin_fidelity"]
    DUR = [30, 70, 70, 60, 70, 60, 70, 70, 70, 115, 60, 70, 105, 70, 60, 60, 60, 30]
    assert sum(DUR) == 1200
    clock = [sum(DUR[:i]) for i in range(len(DUR))]
    pres = new_deck()
    n = 0

    def nxt():
        nonlocal n
        n += 1
        return n - 1

    # ------------------------------------------------------------------ 1 title
    i = nxt()
    s = add_slide(pres, None, dark=True)
    kicker(s, "Happiest Health Digital Twin Challenge 2026", dark=True)
    text(s, MX, 1.75, 6.3, 1.2, "GlucoTwin India", size=60, bold=True, color=IVORY, font=HEAD, name="Deck title")
    text(s, MX, 3.15, 6.0, 1.9,
         f"A personalised digital twin that warns a median of {ev['median_lead_min']:.0f} minutes before a glucose spike, "
         "and lets a doctor test diet and activity changes on the patient's virtual replica before prescribing them.",
         size=20, color=SAND)
    text(s, MX, 5.55, 6.2, 0.4, "Team Vicimus · Vishwakarma Institute of Technology, Pune", size=16, bold=True, color=IVORY)
    text(s, MX, 5.95, 6.2, 0.35, "Chitrangad Ram Sapate (team leader) · Soham Joshi", size=14, color=SAND)
    _, pw, ph = picture(s, ROOT / "docs/src/pptx_assets/hero_crop.png", 7.15, 2.05, w=5.58,
                        alt="The model's output for one synthetic patient: observed glucose, ML forecast with 90 percent interval, and the twin's physiology curve")
    text(s, 7.15, 2.05 + ph + 0.12, 5.58, 0.5, "Real model output for a synthetic patient: observed CGM, the ML forecast with its 90% "
         "interval, and the physiological twin.", size=12, color=SAND)
    notes(s, clock[i], DUR[i], f"""
Good morning. We are Team Vicimus from Vishwakarma Institute of Technology, Pune. This is GlucoTwin India, a digital twin for Type 2 diabetes.

It warns a median of {ev['median_lead_min']:.0f} minutes before a glucose spike, on patients it has never seen, and lets a doctor test a meal or a walk on the patient's virtual copy before prescribing. In twenty minutes: the problem, how it works, the results, and a live demo.

[Timing guide: problem and solution 0:30 to 2:50 · how it works 2:50 to 9:30 · results 9:30 to 13:35 · live demo 13:35 to 16:30 · India, safeguards, limits 16:30 to 19:30 · close 19:30 to 20:00.]""")

    # ------------------------------------------------------------------ 2 problem
    i = nxt()
    s = add_slide(pres, "India manages diabetes by looking in the rear-view mirror")
    kicker(s, "The problem")
    cards = [("101 M", VERM, "Indians living with diabetes, plus 136 M with prediabetes (ICMR-INDIAB, 2023)"),
             ("< 1 in 3", GOLD, "reach glycaemic targets; most complications come from uncontrolled excursions"),
             ("~26,000", ORCHID, "CGM readings per patient between two 3-monthly OPD visits, before counting smartwatch data")]
    cw = (CW - 2 * 0.25) / 3
    for k, (big, col, lab) in enumerate(cards):
        x = MX + k * (cw + 0.25)
        box(s, x, 1.7, cw, 1.95, fill=TINT)
        text(s, x + 0.3, 1.85, cw - 0.6, 0.9, big, size=44, bold=True, color=col, font=HEAD)
        text(s, x + 0.3, 2.75, cw - 0.6, 0.85, lab, size=14, color=TEXT)
    rows = [("HbA1c is an average.", "It hides the post-meal spikes that carbohydrate-heavy diets (rice, wheat) cause every day."),
            ("Care is reactive.", "The doctor sees what already happened, never what is about to happen."),
            ("Advice is generic.", "\"Eat less rice\" ignores that two patients respond to the same meal in completely different ways.")]
    for k, (h1, b1) in enumerate(rows):
        x = MX + k * (cw + 0.25)
        text(s, x, 3.95, cw, 1.5, [{"t": h1, "bold": True, "size": 18, "font": HEAD, "after": 4}, {"t": b1, "size": 14}])
    box(s, MX, 5.6, CW, 1.0, fill=INK)
    text(s, MX + 0.35, 5.6, CW - 0.7, 1.0,
         [{"runs": [("The challenge: ", {"bold": True, "color": IVORY}),
                    ("fuse the static health record with the dynamic wearable stream into a model that predicts adverse events "
                     "before they happen, and that a doctor can actually use.", {"color": IVORY})]}], size=16, anchor="m")
    footer(s, 2, "Sources: ICMR-INDIAB (Lancet Diabetes Endocrinol 2023). 288 CGM readings a day over 90 days is about 26,000.")
    notes(s, clock[i], DUR[i], """
India has about 101 million people with diabetes and another 136 million with prediabetes, according to the ICMR-INDIAB study. Fewer than one in three reach their glycaemic targets.

The reason is how care works. A patient meets the doctor every three months. The doctor looks at HbA1c, which is a single average. Between those two visits a continuous glucose monitor produces about twenty-six thousand readings, and no doctor can read them.

Three consequences. HbA1c is an average, so it hides the post-meal spikes that rice and wheat heavy diets cause every day. Care is reactive: the doctor sees what already happened. And advice is generic: "eat less rice" ignores that two patients can respond to the same meal in completely different ways.

The challenge asks us to fuse the static health record with the dynamic wearable stream, and predict an adverse event before it happens, in a way a doctor can use. That is what we built.""")

    # ------------------------------------------------------------------ 3 solution
    i = nxt()
    s = add_slide(pres, "A virtual patient the doctor can query, not just a risk score")
    kicker(s, "Our solution")
    cards = [("EARLY WARNING", VERM, "\"A spike is likely within 2 hours\"",
              "Forecasts glucose at 30, 60 and 120 minutes with a calibrated 90% interval, and alerts early enough to act: a walk or a smaller portion."),
             ("WHAT-IF", ORCHID, "\"What if she eats jowar roti instead?\"",
              "The doctor simulates meals, portions and post-meal walks on this patient's own calibrated physiology and sees the curve before prescribing."),
             ("EXPLAINED", PERI, "\"Why is the risk high?\"",
              "Every alert lists its patient-specific drivers, and a pre-consultation summary is drafted automatically for the doctor to verify.")]
    for k, (tag, col, q, body) in enumerate(cards):
        x = MX + k * (cw + 0.25)
        box(s, x, 1.7, cw, 3.45, fill=TINT)
        text(s, x + 0.3, 1.95, cw - 0.6, 0.3, tag, size=12, bold=True, color=col)
        text(s, x + 0.3, 2.35, cw - 0.6, 1.0, q, size=20, bold=True, font=HEAD)
        text(s, x + 0.3, 3.5, cw - 0.6, 1.6, body, size=14)
    box(s, MX, 5.4, CW / 2 - 0.125, 1.25, fill=WHITE, line=LINE)
    text(s, MX + 0.3, 5.5, CW / 2 - 0.7, 1.05, [{"t": "Adverse event predicted", "bold": True, "size": 14, "color": VERM, "after": 3},
         {"t": "Glucose above 180 mg/dL within the next 2 hours, with the glucose trajectory at 30, 60 and 120 minutes.", "size": 14}])
    box(s, MX + CW / 2 + 0.125, 5.4, CW / 2 - 0.125, 1.25, fill=WHITE, line=LINE)
    text(s, MX + CW / 2 + 0.425, 5.5, CW / 2 - 0.7, 1.05, [{"t": "Who uses it", "bold": True, "size": 14, "color": PERI, "after": 3},
         {"t": "Endocrinologists and GPs in OPD and tele-consults, and diabetes educators for diet counselling. Next: patients and ASHA workers.", "size": 14}])
    footer(s, 3)
    notes(s, clock[i], DUR[i], f"""
GlucoTwin turns that stream of data into a virtual patient the doctor can interact with, not just a risk score.

Three things. First, early warning. The model forecasts glucose at thirty, sixty and one hundred twenty minutes, with a calibrated ninety percent interval, and flags a spike early enough to act on it, a walk after lunch or a smaller portion.

Second, what-if. The doctor can ask: what happens if she eats jowar roti instead of white rice, or walks fifteen minutes after lunch? The answer comes from that patient's own calibrated physiology.

Third, explained. Every alert comes with its drivers, and a pre-consultation summary is drafted for the doctor to verify.

The adverse event we predict is glucose rising above 180 within the next two hours. The users are endocrinologists and GPs in OPD and tele-consults, and diabetes educators. The next step is patients and ASHA workers.""")

    # ------------------------------------------------------------------ 4 twin properties
    i = nxt()
    s = add_slide(pres, "Five properties that separate a twin from a predictor")
    kicker(s, "Why this is a digital twin")
    props = [("Personalised", PERI, "Six physiological parameters are fitted to each patient, starting from a prior built from their own health record."),
             ("Synchronised", GREEN, "A state observer pulls the twin toward every new CGM reading, so the twin is always \"now\"."),
             ("Predictive", GOLD, "A machine-learning layer learns what physiology leaves out: sleep debt, stress, the dawn effect."),
             ("Simulatable", ORCHID, "Because the core is mechanistic, it answers counterfactuals that a pure ML model cannot."),
             ("Actionable", VERM, "Triage, explanations and summaries sit inside the doctor's workflow.")]
    pw_ = 2.25
    gap = (CW - 5 * pw_) / 4
    for k, (h1, col, b1) in enumerate(props):
        x = MX + k * (pw_ + gap)
        badge(s, x, 1.75, 0.62, str(k + 1), col, size=20)
        text(s, x, 2.55, pw_, 0.45, h1, size=20, bold=True, font=HEAD)
        text(s, x, 3.1, pw_, 2.2, b1, size=15)
    box(s, MX, 5.45, CW, 1.2, fill=INK)
    text(s, MX + 0.35, 5.45, CW - 0.7, 1.2,
         [{"runs": [("Static EHR stream  ×  dynamic wearable stream  →  a physiological replica, updated every 5 minutes.", {"bold": True, "color": IVORY})], "size": 17, "after": 4},
          {"t": "Static: demographics, diagnoses, HbA1c history, labs, medications, TCF7L2 genotype, polygenic risk.   "
                "Dynamic: CGM, heart rate, HRV, steps, sleep stages, meal log.", "size": 14, "color": SAND}], anchor="m")
    footer(s, 4)
    notes(s, clock[i], DUR[i], """
Many projects call a predictor a digital twin. We hold ourselves to five properties.

One, personalised. Six physiological parameters, such as insulin action, glucose effectiveness and how fast this patient absorbs carbohydrate, are fitted to that patient, starting from a prior built from their own health record.

Two, synchronised. A state observer nudges the twin toward every new CGM reading, so the twin is always at the present moment.

Three, predictive. The physiology model is deliberately simple. A machine-learning layer learns what physiology leaves out, like sleep debt, stress and the dawn effect, from the wearable and health-record data.

Four, simulatable. Because the core is mechanistic, it can answer what-if questions that a pure machine-learning model cannot.

Five, actionable. Triage, explanations and summaries sit in the doctor's workflow.

In one line: the static record times the dynamic wearable stream gives a physiological replica, updated every five minutes.""")

    # ------------------------------------------------------------------ 5 data
    i = nxt()
    s = add_slide(pres, "Two data streams fused: the record and the wearable")
    kicker(s, "Data and the sandbox rules")
    colw = (CW - 2 * 0.2) / 3
    cols = [("STATIC STREAM · health record", PERI,
             ["Demographics and diagnoses (SNOMED CT)", "HbA1c history, fasting glucose, lipids, eGFR, blood pressure (LOINC)",
              "Medications: metformin, sulfonylurea, insulin (RxNorm)", "Genetics: TCF7L2 rs7903146, polygenic risk score",
              "Exported as HL7 FHIR R4 bundles with an ABHA identifier"]),
            ("DYNAMIC STREAM · wearables", GREEN,
             ["Continuous glucose monitor, every 5 minutes", "Smartwatch: heart rate, HRV (RMSSD), steps", "Sleep stages: awake, light, deep, REM",
              "Meal log from a 27-dish Indian food library (carbs, GI, fat, protein, fibre)", "Apple Health / Google Fit style, replayed as a live stream"]),
            ("SANDBOX · synthetic and open data only", ORCHID,
             ["80 synthetic Indian T2D and prediabetes patients, 14 days each (seed 2026)", "Our own generator: extended Bergman model with dawn effect, sleep debt, stress, drug action",
              "CGM sensor model: lag, noise, dropouts; imperfect meal logging", "Real-data check on the open CGMacros dataset (PhysioNet)",
              "No real patient data: DPDP Act 2023 and HIPAA rules respected"])]
    for k, (h1, col, items) in enumerate(cols):
        x = MX + k * (colw + 0.2)
        box(s, x, 1.7, colw, 4.85, fill=TINT)
        text(s, x + 0.28, 1.9, colw - 0.56, 0.3, h1, size=12, bold=True, color=col)
        text(s, x + 0.28, 2.4, colw - 0.56, 4.0, [{"t": it, "bullet": True, "after": 10, "size": 15} for it in items])
    footer(s, 5)
    notes(s, clock[i], DUR[i], """
The challenge requires fusing two streams, and the sandbox rules require synthetic or open data.

On the left, the static stream: a simulated health record. Demographics, diagnoses coded in SNOMED, lab results coded in LOINC, medications coded in RxNorm, and genetic markers, specifically the TCF7L2 diabetes risk variant and a polygenic risk score. We export it as HL7 FHIR R4 bundles with an ABHA identifier, so it looks like what a hospital system would send.

In the middle, the dynamic stream: CGM every five minutes, heart rate, heart-rate variability, steps, sleep stages and a meal log from a twenty-seven dish Indian food library.

On the right, how we respect the sandbox rules. All patients are synthetic: eighty Indian patients, fourteen days each, from our own generator that simulates the physiology, including dawn effect, sleep debt, stress and drug action, and a realistic CGM sensor with lag, noise and dropouts. For real-world evidence we validate on the open CGMacros dataset from PhysioNet. No real patient data is used anywhere.""")

    # ------------------------------------------------------------------ 6 architecture
    i = nxt()
    s = add_slide(pres, "Architecture: from raw streams to the doctor's screen")
    kicker(s, "Architecture")
    _, aw, ah = picture(s, ROOT / "docs/architecture.png", MX, 1.65, h=5.0,
                        alt="System architecture diagram: data sources, ingestion and fusion, the twin core with personalise, synchronise, predict and simulate, and the clinical application")
    x0 = MX + aw + 0.45
    stages = [("Personalise", "EHR-informed prior, then MAP calibration of six parameters on days 1 to 10", PERI),
              ("Synchronise", "State observer corrects the twin at every CGM reading", GREEN),
              ("Predict", "LightGBM forecasts with conformal intervals and TreeSHAP reasons", GOLD),
              ("Simulate", "Roll the twin forward 5 hours under planned meals and walks", ORCHID),
              ("Act", "FastAPI service and the clinician dashboard", VERM)]
    for k, (h1, b1, col) in enumerate(stages):
        y = 1.7 + k * 0.99
        badge(s, x0, y + 0.02, 0.42, str(k + 1), col, size=14)
        text(s, x0 + 0.6, y, W - MX - x0 - 0.6, 0.9, [{"t": h1, "bold": True, "size": 16, "font": HEAD, "after": 1}, {"t": b1, "size": 14}])
    footer(s, 6, "Full diagram: docs/architecture.pdf in the repository.")
    notes(s, clock[i], DUR[i], """
This is the system architecture. Read it left to right.

Data sources on the left: the static record and the dynamic wearable stream. Next, ingestion and fusion: the health record becomes a feature vector and seeds the twin's prior, and the wearable stream is aligned to a five-minute grid. A rule we enforce everywhere: every feature uses only data available at that moment.

The twin core in the middle has four steps. Personalise: calibrate six parameters to the patient. Synchronise: a state observer corrects the twin at every reading. Predict: the hybrid machine-learning layer. And Simulate: the what-if engine.

On the right, a FastAPI service with a clinical service layer, and the clinician dashboard, which I will demonstrate shortly.

The full diagram is in the repository as a PDF.""")

    # ------------------------------------------------------------------ 7 twin core
    i = nxt()
    s = add_slide(pres, "A personal physiology model, kept in sync with the patient")
    kicker(s, "Twin core")
    box(s, MX, 1.7, 6.2, 2.15, fill=INK)
    text(s, MX + 0.3, 1.82, 5.6, 1.95,
         [{"t": "G′ = −(SG + X + A)·G + SG·Gb + Ra(t)/V", "after": 6}, {"t": "X′ = p2·(kx·max(G − Gb, 0) − X)", "after": 6},
          {"t": "A′ = (k_ex·cadence − A)/15", "after": 6}, {"t": "Ra = Σ meals f·carbs·s/τ²·e^(−s/τ)"}],
         size=15, color=IVORY, font=MONO, anchor="m", name="Model equations")
    text(s, MX, 4.05, 6.2, 2.7, [
        {"runs": [("EHR prior: ", {"bold": True}), ("basal glucose from fasting glucose, insulin action from HbA1c and disease status", {})], "bullet": True, "after": 7},
        {"runs": [("MAP calibration ", {"bold": True}), ("of six parameters on each patient's days 1 to 10", {})], "bullet": True, "after": 7},
        {"runs": [("Observer: ", {"bold": True}), ("corrects the state at every CGM reading", {})], "bullet": True, "after": 7},
        {"runs": [("Meal absorption ", {"bold": True}), ("from carbs, glycaemic index, fat, protein and fibre", {})], "bullet": True}], size=14)
    _, tw, th = picture(s, FIG / "twin_recovery.png", 7.1, 1.7, w=5.63,
                        alt="Scatter plots of the twin's estimated insulin action and basal glucose against the hidden true values for each synthetic patient")
    y2 = 1.7 + th + 0.22
    sw_ = (5.63 - 0.2) / 2
    for k, (big, lab, col) in enumerate([(f"ρ = {TW['spearman_kx']:.2f}", "insulin action recovered", PERI), (f"ρ = {TW['spearman_gb']:.2f}", "basal glucose recovered", GREEN)]):
        x = 7.1 + k * (sw_ + 0.2)
        box(s, x, y2, sw_, 1.2, fill=TINT)
        text(s, x + 0.25, y2 + 0.12, sw_ - 0.5, 0.6, big, size=28, bold=True, color=col, font=HEAD)
        text(s, x + 0.25, y2 + 0.74, sw_ - 0.5, 0.4, lab, size=13)
    text(s, 7.1, y2 + 1.35, 5.63, 0.7, f"Personal calibration cuts the twin's forecast loss by {TW['mean_loss_reduction_pct']:.0f}% versus the health-record prior alone. "
         "Spearman ρ compares fitted parameters with the hidden ground truth of each synthetic patient.", size=12, color=MUTED)
    footer(s, 7)
    notes(s, clock[i], DUR[i], f"""
The core is a classic physiological model: an oral glucose minimal model, the Bergman model, on a five-minute grid. Glucose G is pulled toward the patient's basal level Gb, reduced by insulin action X and by exercise A, and pushed up by meal absorption Ra.

Three steps make it a twin and not a generic model. First, the health record gives a prior: basal glucose from the fasting glucose, insulin action from HbA1c and disease status. Second, calibration: we fit six parameters to each patient's own first ten days, balancing forecast error against the prior. Third, synchronisation: the observer corrects the state at every reading, so forecasts and simulations start from the patient's current state.

Meal absorption speed comes from each dish's carbohydrate, glycaemic index, fat, protein and fibre.

Does it work? Because our patients are synthetic, we know their hidden true physiology. On the right, the twin's estimates against the truth: a rank correlation of {TW['spearman_kx']:.2f} for insulin action and {TW['spearman_gb']:.2f} for basal glucose. Personal calibration cuts forecast loss by about {TW['mean_loss_reduction_pct']:.0f} percent compared with the record-only prior.""")

    # ------------------------------------------------------------------ 8 ML layer
    i = nxt()
    s = add_slide(pres, "Machine learning corrects what physiology cannot see")
    kicker(s, "Hybrid ML layer")
    g = F["groups"]
    total = sum(g.values())
    cd = CategoryChartData()
    cd.categories = ["CGM history", "Wearables", "Meal log", "Health record", "Twin physiology"]
    cd.add_series("Features", [g["CGM"], g["Wearables"], g["Meal log"], g["EHR"], g["Twin"]])
    gf = s.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(MX), Inches(1.65), Inches(5.6), Inches(4.5), cd)
    ch = gf.chart
    style_chart_text(ch, 13)
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(12)
    ch.has_title = False
    plot = ch.plots[0]
    plot.has_data_labels = True
    plot.data_labels.show_value = True
    plot.data_labels.show_category_name = False
    plot.data_labels.show_percentage = False
    plot.data_labels.font.size = Pt(14)
    plot.data_labels.font.bold = True
    plot.data_labels.font.color.rgb = rgb(WHITE)
    for k, col in enumerate((TEXT, PERI, GOLD, GREEN, ORCHID)):
        pt = plot.series[0].points[k]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = rgb(col)
    gf.name = "Feature groups chart"
    text(s, MX, 6.2, 5.6, 0.5, f"{total} fused features per 5-minute step, by data stream", size=13, color=MUTED, align="c")
    text(s, 6.6, 1.7, 6.13, 3.6, [
        {"runs": [("LightGBM ", {"bold": True}), ("forecasts the change in glucose at 30, 60 and 120 minutes", {})], "bullet": True, "after": 9},
        {"runs": [("Conformalised quantile regression ", {"bold": True}), (f"gives 90% intervals calibrated on held-out patients (measured coverage {f['60']['coverage_pct']:.1f}% at 60 minutes)", {})], "bullet": True, "after": 9},
        {"runs": [("Event classifier ", {"bold": True}), ("gives the probability that glucose exceeds 180 mg/dL within 2 hours", {})], "bullet": True, "after": 9},
        {"runs": [("Twin-simulated trajectories, carbs on board and insulin action ", {"bold": True}), ("are inputs, so physiology and learning inform each other", {})], "bullet": True}], size=15)
    box(s, 6.6, 5.0, 6.13, 1.6, fill=TINT)
    text(s, 6.9, 5.1, 5.55, 1.4, [{"t": "What we do not claim", "bold": True, "size": 14, "color": VERM, "after": 3},
         {"t": "No hypoglycaemia model was trained: the synthetic cohort had too few low-glucose events to learn one, so the dashboard shows no hypo risk rather than an invented number.", "size": 14}])
    footer(s, 8)
    notes(s, clock[i], DUR[i], f"""
The machine-learning layer learns what the physiology model leaves out.

For every five-minute step we fuse {total} features, shown on the left by data stream: CGM trends and variability, wearable signals such as steps, heart rate, HRV and last night's sleep, the meal log, the health record, and the twin's own outputs, simulated trajectories, carbs still absorbing and insulin action.

LightGBM predicts the change in glucose at thirty, sixty and one hundred twenty minutes. For uncertainty we use conformalised quantile regression, calibrated on held-out patients. At sixty minutes, the ninety percent interval contained the true value {f['60']['coverage_pct']:.1f} percent of the time. A classifier gives the probability of exceeding 180 within two hours.

One honest point. We did not train a hypoglycaemia model: the cohort had too few low-glucose events. So the dashboard shows no hypo risk at all, rather than an invented number.""")

    # ------------------------------------------------------------------ 9 alerts + explanation
    i = nxt()
    s = add_slide(pres, "Alerts a clinician can trust: few, early and explained")
    kicker(s, "Alert policy and explainability")
    thr = d["risk"]["spike_threshold"]
    steps_ = [("1", "Pick the operating point the way a clinician would",
               f"Threshold {thr:.3f}: the most sensitive setting with at most 1 false alert per patient-day on calibration patients, not the one that maximises a statistic."),
              ("2", "Confirm before notifying", "Risk must stay above the threshold for two readings (10 minutes), which removes single-reading blips."),
              ("3", "Snooze like a CGM alarm", "No repeat notification within 60 minutes, because alert fatigue is a patient-safety problem.")]
    for k, (nn, h1, b1) in enumerate(steps_):
        y = 1.7 + k * 1.62
        box(s, MX, y, 5.9, 1.48, fill=TINT)
        badge(s, MX + 0.25, y + 0.25, 0.5, nn, PERI, size=16)
        text(s, MX + 0.95, y + 0.17, 4.7, 1.3, [{"t": h1, "bold": True, "size": 15, "after": 3}, {"t": b1, "size": 13}])
    reasons = d["risk"]["reasons"][:5]
    cd = CategoryChartData()
    cd.categories = [x["label"] for x in reasons]
    cd.add_series("Contribution (log-odds)", [x["impact"] for x in reasons])
    text(s, 6.9, 1.62, 5.83, 0.5, [{"t": "Why this patient is flagged", "bold": True, "size": 16, "font": HEAD}])
    gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(6.9), Inches(2.05), Inches(5.83), Inches(3.55), cd)
    gf.name = "TreeSHAP drivers chart"
    ch = gf.chart
    style_chart_text(ch, 12)
    ch.has_legend = False
    ch.has_title = False
    ch.value_axis.visible = False
    ch.value_axis.has_major_gridlines = False
    ch.category_axis.reverse_order = True
    ch.category_axis.tick_labels.font.size = Pt(12)
    ch.category_axis.format.line.color.rgb = rgb(LINE)
    from pptx.enum.chart import XL_TICK_LABEL_POSITION
    ch.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    plot = ch.plots[0]
    plot.gap_width = 45
    plot.has_data_labels = True
    plot.data_labels.number_format = '+0.00;-0.00'
    plot.data_labels.number_format_is_linked = False
    plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    plot.data_labels.font.size = Pt(12)
    ser = plot.series[0]
    ser.invert_if_negative = False
    for k, x in enumerate(reasons):
        pt = ser.points[k]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = rgb(VERM if x["impact"] > 0 else PERI)
    for dpt in ser._element.findall(qn("c:dPt")):      # PowerPoint inverts negative bars (white fill) unless each point says not to
        if dpt.find(qn("c:invertIfNegative")) is None:
            inv = etree.Element(qn("c:invertIfNegative"))
            inv.set("val", "0")
            dpt.find(qn("c:idx")).addnext(inv)
    text(s, 6.9, 5.7, 5.83, 0.95,
         f"{d['patient']['name']} at {DEMO_AT[11:]} on the replay clock: glucose {d['current']['glucose']:.0f} mg/dL, 2-hour risk {d['risk']['p_spike'] * 100:.0f}%. "
         "Red bars raise the risk, blue bars lower it (TreeSHAP, log-odds).", size=12, color=MUTED)
    footer(s, 9)
    notes(s, clock[i], DUR[i], f"""
A model that fires constantly gets switched off, so the alert policy is as important as the model.

Step one: the threshold. We do not maximise a statistic. We choose the most sensitive threshold, {thr:.3f}, that keeps false alerts to at most one per patient per day on calibration patients the model did not train on. Step two: the risk must stay above the threshold for two consecutive readings, ten minutes, before anyone is notified. Step three: after an alert, a sixty-minute snooze, like a CGM alarm.

On the right, the explanation for one alert, which you will see again in the demo. {d['patient']['name']}, glucose {d['current']['glucose']:.0f}, comfortably in range, but a two-hour risk of {d['risk']['p_spike'] * 100:.0f} percent. The bars are TreeSHAP contributions. Current glucose and the time of day push the risk up; the twin's own trajectory pulls it down. The doctor can see why, not just that.""")

    # ------------------------------------------------------------------ 10 results
    i = nxt()
    s = add_slide(pres, "Every data stream adds value, and the hybrid twin is best")
    kicker(s, f"Results · {F['n_ward']} unseen patients, unseen days")
    kp = [(f"{f['60']['rmse']:.1f}", "mg/dL", f"60-minute forecast error (persistence: {pers['60']['rmse']:.1f})", PERI),
          (f"{ev['detected']}/{ev['excursions']}", f"{ev['sensitivity_pct']:.0f}%", f"excursions above 180 flagged in advance, median lead {ev['median_lead_min']:.0f} min", GREEN),
          (f"{ev['false_alerts_per_patient_day']:.2f}", "per day", "false alerts per patient-day (target: at most 1)", GOLD),
          (f"{f['spike']['auroc']:.3f}", "AUROC", f"2-hour spike alert · precision {f['spike']['precision']:.2f} · recall {f['spike']['recall']:.2f}", ORCHID)]
    kw = (CW - 3 * 0.18) / 4
    for k, (big, unit, lab, col) in enumerate(kp):
        x = MX + k * (kw + 0.18)
        box(s, x, 1.65, kw, 1.6, fill=TINT)
        text(s, x + 0.22, 1.72, kw - 0.44, 0.7, [{"runs": [(big, {"size": 30, "bold": True, "color": col, "font": HEAD}), ("  " + unit, {"size": 13, "color": MUTED})]}])
        text(s, x + 0.22, 2.45, kw - 0.44, 0.78, lab, size=12)
    names = ["Persistence", "Linear extrapolation", "Twin only (physiology)", "CGM only", "+ Wearables & meal log", "+ EHR (stream fusion)", "+ Twin physiology (full hybrid)"]
    src = {**m["baselines"], **m["ablation"]}
    cd = CategoryChartData()
    cd.categories = names
    cd.add_series("RMSE at 60 minutes (mg/dL)", [src[nm]["60"]["rmse"] for nm in names])
    gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(MX), Inches(3.45), Inches(CW), Inches(3.25), cd)
    gf.name = "RMSE at 60 minutes by model"
    ch = gf.chart
    style_chart_text(ch, 13)
    ch.has_legend = False
    ch.has_title = True
    ch.chart_title.text_frame.text = "Forecast error at 60 minutes, lower is better (mg/dL)"
    ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(13)
    ch.chart_title.text_frame.paragraphs[0].runs[0].font.bold = True
    ch.value_axis.visible = False
    ch.value_axis.has_major_gridlines = False
    ch.category_axis.reverse_order = True
    ch.category_axis.format.line.color.rgb = rgb(LINE)
    plot = ch.plots[0]
    plot.gap_width = 35
    plot.has_data_labels = True
    plot.data_labels.number_format = "0.0"
    plot.data_labels.number_format_is_linked = False
    plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    plot.data_labels.font.size = Pt(13)
    plot.data_labels.font.bold = True
    for k, nm in enumerate(names):
        pt = plot.series[0].points[k]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = rgb("A8ABBD" if k < 3 else (PERI if k == len(names) - 1 else "8E9CEB"))
    footer(s, 10, f"Protocol: subject split {F['n_train']} / {F['n_cal']} / {F['n_ward']} patients; twins calibrated on days 1 to 10; metrics on days 11 to 14. "
           f"Bars show point-forecast ablation models; the final model with intervals scores {f['60']['rmse']:.1f}.")
    notes(s, clock[i], DUR[i], f"""
Here are the results. The protocol matters: we split by patient, not by row. {F['n_train']} patients train the model, {F['n_cal']} calibrate the uncertainty and alert threshold, and {F['n_ward']} are never seen. Each twin is calibrated on days one to ten only, and every metric is on days eleven to fourteen of the unseen patients.

Top left: the sixty-minute forecast error is {f['60']['rmse']:.1f} milligrams per decilitre, against {pers['60']['rmse']:.1f} for simply assuming glucose stays where it is.

The alerts: of {ev['excursions']} excursions above 180, {ev['detected']} were flagged in advance, {ev['sensitivity_pct']:.0f} percent, with a median warning of {ev['median_lead_min']:.0f} minutes, at {ev['false_alerts_per_patient_day']:.2f} false alerts per patient per day, inside our limit of one. The ranking quality, AUROC, is {f['spike']['auroc']:.3f}.

The chart is the ablation: what each data stream adds. The physiology twin alone beats persistence. CGM alone is better. Adding wearables and meal logs cuts the error again. Adding the health record helps slightly. Adding the twin's physiology gives the best model, {src['+ Twin physiology (full hybrid)']['60']['rmse']:.1f}. Every stream earns its place. The final model with intervals scores {f['60']['rmse']:.1f}.

I should be candid: these are synthetic results, which are an upper bound, because we designed that physiology. That is why the real-data check on the next slides matters.""")

    # ------------------------------------------------------------------ 11 clinical accuracy
    i = nxt()
    s = add_slide(pres, "Clinically safe forecasts, with honest uncertainty")
    kicker(s, "Clinical accuracy")
    rows = [["Horizon", "RMSE (mg/dL)", "MAE", "MARD", "Clarke zone A", "Clarke A + B", "Interval coverage"]]
    for h_ in ("30", "60", "120"):
        x = f[h_]
        rows.append([f"{h_} min", f"{x['rmse']:.1f}", f"{x['mae']:.1f}", f"{x['mard_pct']:.1f}%", f"{x['clarke_A_pct']:.1f}%", f"{x['clarke_AB_pct']:.1f}%", f"{x['coverage_pct']:.1f}%"])
    tbl = s.shapes.add_table(4, 7, Inches(MX), Inches(1.65), Inches(CW), Inches(1.75)).table
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = tbl.cell(ri, ci)
            c.fill.solid()
            c.fill.fore_color.rgb = rgb(INK if ri == 0 else (TINT if ri % 2 else WHITE))
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.margin_left = Inches(0.15)
            tf = c.text_frame
            tf.text = val
            run = tf.paragraphs[0].runs[0]
            run.font.name, run.font.size = BODY, Pt(14)
            run.font.bold = ri == 0 or ci == 0
            run.font.color.rgb = rgb(IVORY if ri == 0 else TEXT)
            tf.paragraphs[0].alignment = PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.RIGHT
    tbl.first_row = True
    tbl.horz_banding = False
    _, cwid, _ = picture(s, FIG / "clarke_60min.png", MX, 3.65, h=3.0, alt="Clarke error grid for the 60-minute forecast; nearly all points fall in zones A and B")
    picture(s, FIG / "forecast_example.png", MX + cwid + 0.3, 3.65, w=CW - cwid - 0.3, h=3.0,
            alt="One unseen test patient-day: observed CGM, the 60-minute-ahead forecast and its 90 percent conformal interval")
    footer(s, 11, "Left: Clarke error grid at 60 minutes. Right: one unseen patient-day, forecast made 60 minutes ahead with its 90% conformal interval.")
    notes(s, clock[i], DUR[i], f"""
Accuracy alone is not safety, so we also report the clinical measures.

The table: mean absolute relative difference, the standard CGM accuracy measure, is {f['60']['mard_pct']:.1f} percent at sixty minutes. The Clarke error grid classifies errors by their clinical consequence. At sixty minutes, {f['60']['clarke_AB_pct']:.1f} percent of forecasts fall in zones A and B, meaning no dangerous treatment decision would follow, and {f['60']['clarke_A_pct']:.1f} percent are in zone A.

On uncertainty: the ninety percent intervals contain the truth {f['30']['coverage_pct']:.0f}, {f['60']['coverage_pct']:.0f} and {f['120']['coverage_pct']:.0f} percent of the time at thirty, sixty and one hundred twenty minutes. Slightly below ninety at the longer horizons, and we report that as it is.

On the right, one unseen patient-day. The line is the forecast made sixty minutes earlier; the band is its interval; the dots are what the sensor actually recorded.""")

    # ------------------------------------------------------------------ 12 real data
    i = nxt()
    s = add_slide(pres, f"It also works on real people: {r['participants']} adults from CGMacros")
    kicker(s, "External validation on real data", color=GREEN)
    order = [("Persistence", rp), ("Twin only (physiology)", r["baselines"]["Twin only (physiology)"]), ("CGM only", r["ablation"]["CGM only"]),
             ("+ Wearable & meal log", r["ablation"]["+ Wearable & meal log"]), ("+ Labs/EHR (stream fusion)", r["ablation"]["+ Labs/EHR (stream fusion)"]),
             ("+ Twin physiology (full hybrid)", rd)]
    rows = [["Model (leave-subjects-out)", "RMSE @30", "@60", "@120"]]
    for nm, sr in order:
        rows.append([nm] + [f"{sr[h_]['rmse']:.1f}" for h_ in ("30", "60", "120")])
    tbl = s.shapes.add_table(len(rows), 4, Inches(MX), Inches(1.65), Inches(7.3), Inches(3.1)).table
    tbl.columns[0].width = Inches(3.7)
    for cc in (1, 2, 3):
        tbl.columns[cc].width = Inches(1.2)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = tbl.cell(ri, ci)
            c.fill.solid()
            last = ri == len(rows) - 1
            c.fill.fore_color.rgb = rgb(INK if ri == 0 else ("E6E9FB" if last else (TINT if ri % 2 else WHITE)))
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.margin_left = Inches(0.15)
            tf = c.text_frame
            tf.text = val
            run = tf.paragraphs[0].runs[0]
            run.font.name, run.font.size = BODY, Pt(14)
            run.font.bold = ri == 0 or last
            run.font.color.rgb = rgb(IVORY if ri == 0 else TEXT)
            tf.paragraphs[0].alignment = PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.RIGHT
    tbl.horz_banding = False
    gain = 100 * (rp["60"]["rmse"] - rd["60"]["rmse"]) / rp["60"]["rmse"]
    kp = [(f"{gain:.0f}%", f"lower 60-minute error than persistence ({rd['60']['rmse']:.1f} vs {rp['60']['rmse']:.1f})", GREEN),
          (f"{re_['detected']}/{re_['excursions']}", f"excursions flagged in advance, median lead {re_['median_lead_min']:.0f} min", PERI),
          (f"{re_['false_alerts_per_day']:.2f}", "false alerts per participant-day: close to, not yet under, the target of 1", GOLD)]
    for k, (big, lab, col) in enumerate(kp):
        y = 1.65 + k * 1.1
        box(s, 8.2, y, 4.53, 1.0, fill=TINT)
        text(s, 8.4, y + 0.1, 1.55, 0.8, big, size=26, bold=True, color=col, font=HEAD, anchor="m")
        text(s, 9.95, y + 0.1, 2.65, 0.8, lab, size=12, anchor="m")
    text(s, MX, 4.95, 7.3, 1.6, [
        {"runs": [("Where it does not help yet: ", {"bold": True, "color": VERM}),
                  (f"the twin alone is worse than persistence at 30 minutes ({r['baselines']['Twin only (physiology)']['30']['rmse']:.1f} vs {rp['30']['rmse']:.1f}); "
                   "fasting labs add no measurable accuracy with 45 people.", {})], "after": 6},
        {"runs": [("Differences from the synthetic study: ", {"bold": True}),
                  ("no HRV or sleep staging, activity from Fitbit METs, a default glycaemic index of 55 for meals, and an American cohort. The same pipeline ran with a single data adapter.", {})]}], size=14)
    text(s, 8.2, 5.0, 4.53, 1.5, f"{sc.get('Healthy', 0)} healthy, {sc.get('Prediabetes', 0)} prediabetes and {sc.get('T2D', 0)} type 2 diabetes participants (by HbA1c). "
         "Twins calibrated on the first 60% of each recording; models never see the test participants; metrics on the last 40%.", size=12, color=MUTED)
    stale = ("\n\n[Presenter note, not spoken: these real-data figures were computed before a later gap-fill fix "
             "that moved the synthetic results by 1 to 3 percent; a re-run is pending.]") if REAL_DATA_STALE else ""
    footer(s, 12, "Data: Gutierrez-Osuna et al., CGMacros v1.0.0, PhysioNet 2025, CC BY-NC-SA 4.0." + ("  Computed before a later gap-fill fix; re-run pending." if REAL_DATA_STALE else ""))
    notes(s, clock[i], DUR[i], f"""
Synthetic results are an upper bound, so we ran the same pipeline on {r['participants']} real adults from the open CGMacros dataset: {sc.get('Healthy', 0)} healthy, {sc.get('Prediabetes', 0)} prediabetic and {sc.get('T2D', 0)} with type 2 diabetes. Twins are calibrated on the first sixty percent of each recording, the models never see the test participants, and metrics are on the final forty percent.

The full hybrid reaches {rd['60']['rmse']:.1f} at sixty minutes against {rp['60']['rmse']:.1f} for persistence, {gain:.0f} percent lower, and beats the CGM-only model at every horizon. Alerts flagged {re_['detected']} of {re_['excursions']} excursions, median lead {re_['median_lead_min']:.0f} minutes, at {re_['false_alerts_per_day']:.2f} false alerts a day: close to our target of one, not yet under it.

Where it does not help: the twin alone is worse than persistence at thirty minutes, and fasting labs add nothing measurable with forty-five people. The cohort is American, with no HRV or sleep data.{stale}""")

    # ------------------------------------------------------------------ 13 doctor's view
    i = nxt()
    s = add_slide(pres, "The doctor's view: from the ward to one patient's twin")
    kicker(s, "Live demo")
    _, w1, h1_ = picture(s, IMG / "dashboard_top.png", MX, 1.65, w=6.5, alt="Overview page: hero chart of the highest-risk patient, four KPIs, risk queue and triage map")
    _, w2, _ = picture(s, FIG / "twin.png", MX + w1 + 0.25, 1.65, h=h1_, alt="Patient twin page: two-hour risk with drivers, observed, forecast and twin series, and the replay scrubber")
    yc = 1.65 + h1_ + 0.12
    text(s, MX, yc, w1, 0.85, [{"t": "Overview", "bold": True, "size": 15, "font": HEAD, "after": 2},
         {"t": "Risk queue with each patient's main driver, a triage map of glucose against risk, and recent alerts with their outcomes.", "size": 13}])
    text(s, MX + w1 + 0.25, yc, w2, 0.85, [{"t": "Patient twin", "bold": True, "size": 15, "font": HEAD, "after": 2},
         {"t": "Risk and its TreeSHAP drivers; observed, forecast and twin series; a replay scrubber over the whole record.", "size": 13}])
    box(s, MX, 6.42, CW, 0.55, fill=INK)
    text(s, MX + 0.3, 6.42, CW - 0.6, 0.55,
         [{"runs": [("Open the live site now:  ", {"color": SAND}), (LIVE_URL.replace("https://", ""), {"bold": True, "color": IVORY}),
                    ("   ·   replay clock at Fri 11 Sept, 12:10", {"color": SAND})]}], size=15, anchor="m")
    footer(s, 13)
    notes(s, clock[i], DUR[i], f"""
[LIVE DEMO. Open {LIVE_URL} before the talk and keep it ready. If it has been idle, the first load can take a few seconds; load it once before you start.]

This is the doctor's view. The header has a replay clock: the dataset is recorded, and we replay it as if live. I have set it to Friday the eleventh of September, twelve-ten.

Overview. Four numbers: patients monitored, how many need attention, alerts in the last day, data freshness. Then the risk queue, ranked by risk tier and the model's two-hour probability. Each row says which factor is driving the risk. [Click the top patient or press 'Open highest-risk twin'.] On the right, the triage map: each dot is a patient, glucose across, risk up; the shaded band is above the alert threshold.

Now the patient twin page, say {d['patient']['name']}. [Click Ward, then P050 or use the link.] Two-hour risk and its drivers on the left. On the chart: observed glucose, the ML forecast with its interval, and the twin's own curve, each drawn and labelled differently so none can pass for another. The strip underneath is this patient's whole record up to the clock. [Drag the scrubber or press the right arrow.] The chart, risk and twin state all follow the clock. [Toggle 'Reveal what happened next'.] Amber circles are what the sensor recorded afterwards. They were never shown to the model. Over those two hours glucose actually peaked at {F['peak']:.0f}, and the forecast for the two-hour mark was {d['forecast']['horizons']['120']['point']:.0f}.""")

    # ------------------------------------------------------------------ 14 what-if
    i = nxt()
    s = add_slide(pres, "Ask the twin before you prescribe")
    kicker(s, "What-if simulator", color=ORCHID)
    _, w1, h1_ = picture(s, IMG / "whatif.png", MX, 1.65, w=7.2, alt="What-if page: plan controls and the simulated glucose curves for no further meal, a planned meal, and the meal with a 15-minute walk")
    sc3 = wi["scenarios"]
    cards = [("No further meal", sc3[0], PERI, None), ("White rice, dal and sabzi", sc3[1], ORCHID, None), ("+ 15-minute walk after the meal", sc3[2], GREEN, sc3[1])]
    for k, (lab, sc_, col, ref) in enumerate(cards):
        y = 1.65 + k * 1.37
        box(s, 8.1, y, 4.63, 1.25, fill=TINT)
        text(s, 8.3, y + 0.08, 4.25, 0.3, lab, size=13, bold=True)
        text(s, 8.3, y + 0.42, 4.25, 0.5, [{"runs": [(f"{sc_['peak']}", {"size": 24, "bold": True, "color": col, "font": HEAD}), (f" mg/dL peak · {sc_['time_above_180_min']} min above 180", {"size": 13, "color": MUTED})]}])
        if ref:
            text(s, 8.3, y + 0.92, 4.25, 0.3, f"{sc_['peak'] - ref['peak']:+d} peak and {sc_['time_above_180_min'] - ref['time_above_180_min']:+d} min above 180 versus the meal alone", size=11, color=GREEN)
    box(s, MX, 5.85, CW, 0.9, fill=WHITE, line=LINE)
    text(s, MX + 0.3, 5.85, CW - 0.6, 0.9, [{"runs": [("Assumptions: ", {"bold": True}),
         ("the twin starts from its state at the replay clock, with logged meals still absorbing; the walk is 100 steps a minute using this patient's calibrated exercise response; "
          "medication, stress and sleep changes are not simulated; the physiology is deterministic, so no uncertainty band is drawn and the result is not a promise.", {})]}], size=13, anchor="m")
    footer(s, 14)
    pk0, pk1, pk2 = sc3[0]["peak"], sc3[1]["peak"], sc3[2]["peak"]
    notes(s, clock[i], DUR[i], f"""
[LIVE DEMO. On the live site click 'Simulate' in the header.]

This is the what-if simulator: the page that makes it a twin rather than a predictor. The twin is rolled forward five hours from the replay clock under three futures: no further meal, a planned lunch of white rice, dal and sabzi, and the same lunch followed by a fifteen-minute walk.

For this patient, with no further meal glucose stays around {pk0}. Eating the lunch sends it to a peak of {pk1}, and {sc3[1]['time_above_180_min']} minutes above 180. Adding a fifteen-minute walk lowers the peak to {pk2}, {pk1 - pk2} points lower, and {sc3[1]['time_above_180_min'] - sc3[2]['time_above_180_min']} fewer minutes above 180. [Move the portion or walk slider and press Simulate. Hover the chart to read all three at the same minute.]

Press 'Rank options' and the twin simulates every region-appropriate alternative for this meal slot for this patient, ranked by time above 180. That is the "what if she eats jowar roti instead" question, answered on her own physiology.

Be clear about limits: this is deterministic physiology. It does not simulate medication, stress or sleep changes, and it shows no uncertainty band, which must not be read as certainty.""")

    # ------------------------------------------------------------------ 15 India
    i = nxt()
    s = add_slide(pres, "Designed around Indian patients, food and health systems")
    kicker(s, "Built for India")
    iw = (CW - 3 * 0.2) / 4
    items = [("Indian diet", GOLD, "27 dishes (idli, dosa, poha, paratha, rajma chawal, biryani, ragi, jowar, chai and biscuits) with glycaemic index and macros driving absorption. North and South diets are modelled."),
             ("South Asian profile", PERI, "Diabetes at lower BMI, higher insulin resistance, and a TCF7L2 risk-allele frequency of about 0.28 in the synthetic cohort."),
             ("ABDM-ready", GREEN, "HL7 FHIR R4 resources with SNOMED CT, LOINC and RxNorm codes and an ABHA identifier, masked in the UI and the API."),
             ("DPDP Act 2023", VERM, "Synthetic data only, purpose-limited fields, masking, and one container that runs on-premise.")]
    for k, (h1, col, b1) in enumerate(items):
        x = MX + k * (iw + 0.2)
        box(s, x, 1.7, iw, 3.1, fill=TINT)
        dot(s, x + 0.28, 1.95, 0.22, col)
        text(s, x + 0.28, 2.3, iw - 0.5, 0.4, h1, size=16, bold=True, font=HEAD)
        text(s, x + 0.28, 2.85, iw - 0.5, 1.9, b1, size=13)
    hw = (CW - 0.2) / 2
    box(s, MX, 5.05, hw, 1.55, fill=WHITE, line=LINE)
    text(s, MX + 0.3, 5.15, hw - 0.6, 1.4, [{"t": "Pre-consultation summary", "bold": True, "size": 15, "font": HEAD, "after": 3},
         {"t": "An auto-drafted narrative (time in range, GMI, variability, food response, sleep effect, twin phenotype) with discussion points, so a 5-minute OPD slot starts informed.", "size": 13}])
    box(s, MX + hw + 0.2, 5.05, hw, 1.55, fill=WHITE, line=LINE)
    text(s, MX + hw + 0.5, 5.15, hw - 0.6, 1.4, [{"t": "A low-cost path to scale", "bold": True, "size": 15, "font": HEAD, "after": 3},
         {"t": "Designed for affordable 14-day CGMs and any smartwatch. Next: Hindi and Kannada alerts and a low-bandwidth mode for ASHA workers.", "size": 13}])
    footer(s, 15)
    notes(s, clock[i], DUR[i], """
We designed this for India, not adapted from a Western product.

Diet: a twenty-seven dish library, idli, dosa, poha, paratha, rajma chawal, biryani, ragi, jowar, each with glycaemic index and macros that drive absorption. North and South diets are modelled separately, and meal swaps are region-appropriate.

Phenotype: South Asians develop diabetes at lower BMI with higher insulin resistance, and our synthetic cohort reflects that.

Infrastructure: FHIR R4 resources with standard codes and an ABHA identifier, masked in the interface and the API, which fits the ABDM ecosystem.

Privacy: synthetic data only, masking, and one on-premise container, in line with the DPDP Act. And a pre-consultation summary, so a five-minute OPD slot starts informed.""")

    # ------------------------------------------------------------------ 16 engineering
    i = nxt()
    s = add_slide(pres, "Built to be trusted: tests, safeguards and open data")
    kicker(s, "Engineering quality")
    gw = (CW - 2 * 0.2) / 3
    cards = [("Leakage-proof by construction", PERI, "Every feature uses only data available at that moment. A unit test perturbs the future, including right after a sensor dropout, and checks that present features do not change."),
             ("The API cannot peek", GREEN, "Tests check that outcomes and alerts use only data up to the replay clock, and that revealing the recorded future never changes the model's risk."),
             ("Honest labelling", ORCHID, "Observed, predicted, simulated and revealed data have different marks and labels everywhere. No hypo risk is shown because no hypo model exists."),
             ("Subject-level evaluation", GOLD, "Patients are split into train, calibration and test sets. Test patients are never seen, and results are on unseen days."),
             ("Privacy and safety", VERM, "ABHA numbers masked in the UI and API, no cross-site access by default, request size limits, and synthetic data only."),
             ("Reproducible and deployed", PERI, "16 automated tests, a one-command pipeline, a Docker image, and a live deployment. Results regenerate from a fixed seed.")]
    for k, (h1, col, b1) in enumerate(cards):
        x = MX + (k % 3) * (gw + 0.2)
        y = 1.7 + (k // 3) * 2.5
        box(s, x, y, gw, 2.3, fill=TINT)
        dot(s, x + 0.28, y + 0.25, 0.2, col)
        text(s, x + 0.28, y + 0.58, gw - 0.5, 0.4, h1, size=16, bold=True, font=HEAD)
        text(s, x + 0.28, y + 1.0, gw - 0.5, 1.25, b1, size=13)
    footer(s, 16)
    notes(s, clock[i], DUR[i], """
A few things that make this trustworthy rather than just impressive.

Leakage. The easiest way to fake a good result is to let the model see the future. Every feature uses only data available at that moment, and a unit test perturbs the future, including right after a sensor dropout. We actually found and fixed a leak of exactly that kind late in the project. It changed our headline numbers slightly, and the numbers you see are after the fix.

The API cannot peek either: tests check that outcomes use only data up to the replay clock, and that revealing the recorded future never changes the model's risk.

Observed, predicted, simulated and revealed data always look different, and we show no hypoglycaemia risk because we have no such model. Evaluation is by patient. And it is reproducible: sixteen tests, a one-command pipeline, a Docker image and a live deployment.""")

    # ------------------------------------------------------------------ 17 limits + roadmap
    i = nxt()
    s = add_slide(pres, "What a proof of concept can and cannot claim")
    kicker(s, "Honest limitations and roadmap")
    hw = (CW - 0.25) / 2
    box(s, MX, 1.7, hw, 4.85, fill=TINT)
    text(s, MX + 0.35, 1.9, hw - 0.7, 0.4, "Limitations", size=20, bold=True, font=HEAD, color=VERM)
    text(s, MX + 0.35, 2.5, hw - 0.7, 4.0, [{"t": t, "bullet": True, "after": 11, "size": 15} for t in [
        "Synthetic results are an upper bound: we designed that physiology.",
        "CGMacros is small (45 people), American, and has no sleep or HRV data; its meals use a default glycaemic index.",
        "No hypoglycaemia model: too few low-glucose events to learn one.",
        "The twin does not model medication changes, illness or stress; the ML layer only partly captures them.",
        "Decision support only. Not a medical device, and not validated on Indian patients yet."]])
    box(s, MX + hw + 0.25, 1.7, hw, 4.85, fill=TINT)
    text(s, MX + hw + 0.6, 1.9, hw - 0.7, 0.4, "Roadmap", size=20, bold=True, font=HEAD, color=PERI)
    text(s, MX + hw + 0.6, 2.5, hw - 0.7, 4.0, [{"runs": [(a, {"bold": True}), (b, {})], "bullet": True, "after": 11, "size": 15} for a, b in [
        ("Validate ", "on Indian CGM cohorts with partner diabetologists."),
        ("Twin upgrades: ", "particle-filter state estimation and medication-titration simulation."),
        ("Patient app: ", "Hindi and Kannada alerts, and WhatsApp or SMS for ASHA workers."),
        ("ABDM sandbox ", "integration with a consent flow, then a prospective clinic pilot."),
        ("More open data: ", "the Shanghai T2DM dataset and larger FHIR cohorts from Synthea.")]])
    footer(s, 17)
    notes(s, clock[i], DUR[i], """
What can a proof of concept claim, and what can it not?

Limitations, plainly. The synthetic results are an upper bound, because we designed the physiology that generated the data. The real-data check is small, American, and has no sleep or heart-rate-variability data. We have no hypoglycaemia model. The twin does not model medication changes, illness or stress. It is decision support, not a medical device, and it has not been validated on Indian patients.

Roadmap. First, validation on Indian CGM cohorts with partner diabetologists. Second, upgrades to the twin: better state estimation, and simulation of medication titration. Third, a patient app with Hindi and Kannada alerts, plus a low-bandwidth mode so ASHA workers can use it by WhatsApp or SMS. Fourth, ABDM sandbox integration and a prospective clinic pilot. And more open data, including the Shanghai dataset and larger FHIR cohorts.""")

    # ------------------------------------------------------------------ 18 close
    i = nxt()
    s = add_slide(pres, None, dark=True)
    kicker(s, "Thank you", dark=True)
    text(s, MX, 1.7, 9.5, 2.0, "From reactive HbA1c reviews to a twin that sees the next two hours.", size=44, bold=True, color=IVORY, font=HEAD, name="Closing statement")
    text(s, MX, 4.2, 9, 0.4, "Team Vicimus · Vishwakarma Institute of Technology, Pune", size=18, bold=True, color=IVORY)
    text(s, MX, 4.65, 9, 0.35, "Chitrangad Ram Sapate (team leader) · Soham Joshi", size=15, color=SAND)
    text(s, MX, 5.5, 11, 1.2, [{"runs": [("Live demo:  ", {"color": SAND}), (LIVE_URL.replace("https://", ""), {"bold": True, "color": IVORY})], "after": 4},
                                {"runs": [("Code (MIT):  ", {"color": SAND}), (REPO_URL, {"bold": True, "color": IVORY})]}], size=16)
    footer(s, 18, dark=True)
    notes(s, clock[i], DUR[i], f"""
To sum up. GlucoTwin India fuses a health record and a wearable stream into a personal twin. It warns a median of {ev['median_lead_min']:.0f} minutes before a spike on unseen patients, explains why, and lets the doctor test a meal or a walk first. We validated it on synthetic Indian patients and {r['participants']} real adults, we are open about its limits, and it is open source under MIT.

Thank you. We would be glad to take your questions.""")

    cp = pres.core_properties
    cp.author = cp.last_modified_by = "Team Vicimus"
    cp.title = "GlucoTwin India: personalised digital twin for Type 2 diabetes"
    cp.subject = "Happiest Health Digital Twin Challenge 2026"
    cp.comments = "20-minute presentation with speaker notes"
    cp.keywords = "digital twin, diabetes, CGM, FHIR, India"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    pres.save(OUT)
    print(f"wrote {OUT.relative_to(ROOT)} ({n} slides, {sum(DUR) // 60} min)")


if __name__ == "__main__":
    build()
